import shutil

import docx
import pytest

from lawproc import config
from lawproc.convert import convert_doc_files

SAMPLE = config.RAW / "legal" / "nd_51_2025__p2.doc"


def test_copies_existing_docx(tmp_path):
    src, out = tmp_path / "src", tmp_path / "out"
    src.mkdir()
    d = docx.Document()
    d.add_paragraph("Điều 1.")
    d.save(src / "a__p1.docx")
    assert [p.name for p in convert_doc_files(src, out)] == ["a__p1.docx"]


@pytest.mark.skipif(not SAMPLE.exists(), reason="chưa có dữ liệu thô")
def test_converts_real_doc_with_word(tmp_path):
    src, out = tmp_path / "src", tmp_path / "out"
    src.mkdir()
    shutil.copy(SAMPLE, src / SAMPLE.name)
    result = convert_doc_files(src, out)
    assert [p.name for p in result] == ["nd_51_2025__p2.docx"]
    text = "\n".join(p.text for p in docx.Document(result[0]).paragraphs)
    assert "10/2022/NĐ-CP" in text
