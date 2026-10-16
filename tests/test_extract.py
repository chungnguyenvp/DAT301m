from pathlib import Path

import docx
import pymupdf
import pytest

from lawproc import config
from lawproc.extract import docx_lines, pdf_lines, is_scanned, ocr_pdf_lines, text_lines

ARIAL = Path("C:/Windows/Fonts/arial.ttf")


def test_text_lines_strips_bom_blank_lines_and_crlf(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes("\ufeffBỘ Y TẾ\r\n\r\n  Điều 1. Phạm vi điều chỉnh  \r\n\t\r\n1. Nội dung\r\n".encode("utf-8"))
    assert text_lines(p) == ["BỘ Y TẾ", "Điều 1. Phạm vi điều chỉnh", "1. Nội dung"]


def test_text_lines_normalizes_to_nfc(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("khỏe".replace("ỏ", "o\u0309"), encoding="utf-8")  # dạng tổ hợp NFD
    assert text_lines(p) == ["khỏe"]


def test_docx_lines_keeps_paragraph_and_table_order(tmp_path):
    d = docx.Document()
    d.add_paragraph("CHÍNH PHỦ")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Số: 168/2024/NĐ-CP"
    t.cell(0, 1).text = "Hà Nội, ngày 26 tháng 12 năm 2024"
    d.add_paragraph("Điều 1. Phạm vi điều chỉnh")
    d.add_paragraph("   ")
    p = tmp_path / "a.docx"
    d.save(p)
    assert docx_lines(p) == [
        "CHÍNH PHỦ",
        "Số: 168/2024/NĐ-CP | Hà Nội, ngày 26 tháng 12 năm 2024",
        "Điều 1. Phạm vi điều chỉnh",
    ]


def test_docx_lines_keeps_breaks_and_tabs_as_spaces(tmp_path):
    d = docx.Document()
    p1 = d.add_paragraph()
    r1 = p1.add_run("VIỆT NAM")
    r1.add_break()
    p1.add_run("Độc lập")
    p2 = d.add_paragraph()
    r2 = p2.add_run("Số:")
    r2.add_tab()
    p2.add_run("168/2024")
    p = tmp_path / "b.docx"
    d.save(p)
    assert docx_lines(p) == ["VIỆT NAM Độc lập", "Số: 168/2024"]


_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _numbered_doc(path, abstract_levels, nums, paragraphs, table_paragraphs=()):
    """Tạo docx có numbering.xml riêng.

    abstract_levels: {abstractNumId: [(numFmt, lvlText, start), ...]} theo ilvl 0, 1, ...
    nums: {numId: (abstractNumId, {ilvl: startOverride})}
    paragraphs / table_paragraphs: [(text, numId | None, ilvl)]
    """
    from docx.oxml import parse_xml

    d = docx.Document()
    numbering = d.part.numbering_part.element
    for child in list(numbering):
        numbering.remove(child)
    for aid, levels in abstract_levels.items():
        lv = "".join(
            f'<w:lvl w:ilvl="{i}"><w:start w:val="{s}"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{txt}"/></w:lvl>'
            for i, (fmt, txt, s) in enumerate(levels))
        numbering.append(parse_xml(f'<w:abstractNum {_W} w:abstractNumId="{aid}">{lv}</w:abstractNum>'))
    for nid, (aid, overrides) in nums.items():
        ov = "".join(f'<w:lvlOverride w:ilvl="{i}"><w:startOverride w:val="{s}"/></w:lvlOverride>'
                     for i, s in overrides.items())
        numbering.append(parse_xml(f'<w:num {_W} w:numId="{nid}"><w:abstractNumId w:val="{aid}"/>{ov}</w:num>'))

    def number(par, nid, ilvl):
        if nid is not None:
            par._p.get_or_add_pPr().append(parse_xml(
                f'<w:numPr {_W}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{nid}"/></w:numPr>'))

    for text, nid, ilvl in paragraphs:
        number(d.add_paragraph(text), nid, ilvl)
    if table_paragraphs:
        cell = d.add_table(rows=1, cols=1).cell(0, 0)
        first = True
        for text, nid, ilvl in table_paragraphs:
            par = cell.paragraphs[0] if first else cell.add_paragraph()
            par.text, first = text, False
            number(par, nid, ilvl)
    d.save(path)


def test_docx_lines_renders_decimal_and_letter_numbering_with_restart(tmp_path):
    p = tmp_path / "n.docx"
    _numbered_doc(p, {"0": [("decimal", "%1.", 1), ("lowerLetter", "%2)", 1)]}, {"1": ("0", {})},
                  [("Điều 3. Giải thích từ ngữ", None, 0), ("Đường đôi", "1", 0), ("Làn đường", "1", 1),
                   ("Phần đường", "1", 1), ("Đường một chiều", "1", 0), ("Xe cơ giới", "1", 1)])
    assert docx_lines(p) == ["Điều 3. Giải thích từ ngữ", "1. Đường đôi", "a) Làn đường", "b) Phần đường",
                             "2. Đường một chiều", "a) Xe cơ giới"]


def test_docx_lines_renders_bullets_overrides_and_multilevel(tmp_path):
    p = tmp_path / "n.docx"
    _numbered_doc(p, {"0": [("bullet", "", 1)], "1": [("decimal", "%1.", 1), ("decimal", "%1.%2.", 1)]},
                  {"1": ("0", {}), "2": ("1", {0: 4}), "3": ("1", {})},
                  [("Gạch đầu dòng", "1", 0), ("Bắt đầu từ bốn", "2", 0), ("Mục con", "2", 1),
                   ("Không đánh số", "0", 0)])
    assert docx_lines(p) == ["- Gạch đầu dòng", "4. Bắt đầu từ bốn", "4.1. Mục con", "Không đánh số"]


def test_docx_lines_continues_lists_sharing_abstract_unless_restarted(tmp_path):
    p = tmp_path / "n.docx"
    _numbered_doc(p, {"0": [("decimal", "%1.", 1)]}, {"1": ("0", {}), "2": ("0", {}), "3": ("0", {0: 1})},
                  [("Một", "1", 0), ("Hai", "2", 0), ("Ba", "1", 0), ("Bắt đầu lại", "3", 0), ("Tiếp", "3", 0)])
    assert docx_lines(p) == ["1. Một", "2. Hai", "3. Ba", "1. Bắt đầu lại", "2. Tiếp"]


def test_docx_lines_keeps_typed_label_instead_of_auto_label(tmp_path):
    p = tmp_path / "n.docx"
    _numbered_doc(p, {"0": [("decimal", "%1.", 1)]}, {"1": ("0", {})},
                  [("Mở đầu", "1", 0), ("1. Nhiệm vụ", "1", 0), ("a) Điểm gõ tay", "1", 0), ("Tiếp", "1", 0)])
    assert docx_lines(p) == ["1. Mở đầu", "1. Nhiệm vụ", "a) Điểm gõ tay", "4. Tiếp"]


def test_docx_lines_numbers_paragraphs_inside_tables(tmp_path):
    p = tmp_path / "n.docx"
    _numbered_doc(p, {"0": [("decimal", "%1.", 1)]}, {"1": ("0", {})},
                  [], table_paragraphs=[("Thắt dây an toàn", "1", 0), ("Khởi hành nhẹ nhàng", "1", 0)])
    assert docx_lines(p) == ["1. Thắt dây an toàn 2. Khởi hành nhẹ nhàng"]


def test_docx_lines_counts_empty_numbered_paragraphs(tmp_path):
    p = tmp_path / "n.docx"
    _numbered_doc(p, {"0": [("decimal", "%1.", 1)]}, {"1": ("0", {})},
                  [("Một", "1", 0), ("", "1", 0), ("Ba", "1", 0)])
    assert docx_lines(p) == ["1. Một", "3. Ba"]


def test_docx_lines_splits_soft_break_before_a_labelled_item(tmp_path):
    d = docx.Document()
    par = d.add_paragraph()
    run = par.add_run("2. Phân loại đường bộ.")
    run.add_break()
    par.add_run("3. Phần đất để bảo vệ đường bộ.")
    p = tmp_path / "br.docx"
    d.save(p)
    assert docx_lines(p) == ["2. Phân loại đường bộ.", "3. Phần đất để bảo vệ đường bộ."]


def _pdf_with_lines(path, lines):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="arial", fontfile=str(ARIAL))
    for k, text in enumerate(lines):
        page.insert_text((72, 72 + 22 * k), text, fontname="arial", fontsize=14)
    doc.save(path)


def test_pdf_lines_and_not_scanned(tmp_path):
    p = tmp_path / "t.pdf"
    _pdf_with_lines(p, ["Điều 1. Phạm vi điều chỉnh của Nghị định này"] * 8)
    assert pdf_lines(p)[0].replace("\xa0", " ").startswith("Điều 1.")  # PyMuPDF may emit NBSP; normalize() cleans it later
    assert is_scanned(p) is False


def test_is_scanned_true_for_image_only_pdf(tmp_path):
    p = tmp_path / "img.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(p)
    assert is_scanned(p) is True


@pytest.mark.skipif(not config.TESSERACT.exists(), reason="Tesseract chưa cài")
def test_ocr_pdf_lines_reads_vietnamese_and_caches(tmp_path):
    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2. Người nộp lệ phí"])
    cache = tmp_path / "cache"
    lines = ocr_pdf_lines(p, cache, dpi=200)
    assert any("Người nộp lệ phí" in x for x in lines)
    assert (cache / "dpi200" / "p0000.txt").exists()
    assert (cache / "source.sha256").read_text().strip()


def test_ocr_cache_from_other_file_is_refused(tmp_path):
    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    cache = tmp_path / "cache"
    cache.mkdir()
    (cache / "source.sha256").write_text("0" * 64 + "\n")
    with pytest.raises(ValueError, match="sha256"):
        ocr_pdf_lines(p, cache, dpi=200)


@pytest.mark.parametrize("other,expected", [("Người nộp lệ phí", "Người nộp lệ phí"),
                                           ("Xe ô tô", "Người nộp lệ phí"),
                                           ("xyz xyz", "Người nộp lệ phí")])
def test_ocr_retry_selection_and_ties(tmp_path, monkeypatch, other, expected):
    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    calls = []

    def fake_page(page, i, cache_dir, dpi, tmp):
        calls.append((i, dpi))
        return "Người nộp lệ phí" if dpi == 200 else other

    monkeypatch.setattr("lawproc.extract._ocr_page", fake_page)
    assert ocr_pdf_lines(p, tmp_path / "cache", dpi=200, retry_pages=[0], retry_dpi=400) == [expected]
    assert calls == [(0, 200), (0, 400)]


def test_ocr_retry_pages_keep_better_text(tmp_path, monkeypatch):
    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    texts = {200: "la Xeôtô wW0W0W lbXxemt\n", 400: "Người nộp lệ phí\n"}
    monkeypatch.setattr("lawproc.extract._ocr_page", lambda page, i, cache_dir, dpi, tmp: texts[dpi])
    cache = tmp_path / "cache"
    assert ocr_pdf_lines(p, cache, dpi=200, retry_pages=[0]) == ["Người nộp lệ phí"]
    assert ocr_pdf_lines(p, cache, dpi=200) == ["la Xeôtô wW0W0W lbXxemt"]


def test_ocr_empty_retry_does_not_replace_nonempty_page(tmp_path, monkeypatch):
    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    monkeypatch.setattr("lawproc.extract._ocr_page", lambda page, i, cache_dir, dpi, tmp:
                        "xyz xyz" if dpi == 200 else " \n")
    assert ocr_pdf_lines(p, tmp_path / "cache", dpi=200, retry_pages=[0]) == ["xyz xyz"]


@pytest.mark.parametrize("pages", [[-1], [1], [0.0], [True], [0, False]])
def test_ocr_invalid_retry_pages_are_refused(tmp_path, pages):
    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    with pytest.raises(ValueError, match="Trang OCR"):
        ocr_pdf_lines(p, tmp_path / "cache", retry_pages=pages)


@pytest.mark.parametrize("dpi", [0, -1, 200.5, True])
def test_ocr_invalid_dpi_is_refused(tmp_path, dpi):
    with pytest.raises(ValueError, match="DPI"):
        ocr_pdf_lines(tmp_path / "scan.pdf", tmp_path / "cache", dpi=dpi)


def test_ocr_unkeyed_cache_is_refused(tmp_path):
    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    cache = tmp_path / "cache"
    cache.mkdir()
    (cache / "p0000.txt").write_text("Nguồn chưa được xác minh", encoding="utf-8")
    with pytest.raises(ValueError, match="thiếu source.sha256"):
        ocr_pdf_lines(p, cache)


def test_ocr_invalid_utf8_cache_is_refused(tmp_path):
    from lawproc.extract import _check_source

    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    cache = tmp_path / "cache"
    _check_source(p, cache)
    (cache / "dpi300").mkdir()
    (cache / "dpi300" / "p0000.txt").write_bytes(b"\xff")
    with pytest.raises(ValueError, match="UTF-8"):
        ocr_pdf_lines(p, cache)


def test_ocr_failure_does_not_cache_page(tmp_path, monkeypatch):
    from types import SimpleNamespace

    p = tmp_path / "scan.pdf"
    _pdf_with_lines(p, ["Điều 2."])
    monkeypatch.setattr("lawproc.extract.subprocess.run", lambda *args, **kwargs:
                        SimpleNamespace(returncode=1, stderr="lỗi thử nghiệm", stdout=""))
    cache = tmp_path / "cache"
    with pytest.raises(RuntimeError, match="trang 0"):
        ocr_pdf_lines(p, cache, dpi=72)
    assert not (cache / "dpi72" / "p0000.txt").exists()


def test_docx_lines_reads_nested_table_once(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Mẫu số 01"
    inner = t.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "Họ và tên"
    inner.cell(0, 1).text = "Ngày sinh"
    p = tmp_path / "n.docx"
    d.save(p)
    assert docx_lines(p) == ["Mẫu số 01 | Họ và tên | Ngày sinh"]


def test_docx_lines_merged_cells_once(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Tiêu đề gộp"
    t.cell(1, 0).text = "Trái"
    t.cell(1, 1).text = "Phải"
    p = tmp_path / "m.docx"
    d.save(p)
    assert docx_lines(p) == ["Tiêu đề gộp", "Trái | Phải"]


def test_docx_lines_cell_with_empty_paragraph_has_single_spaces(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    c.text = "Dòng một"
    c.add_paragraph("")
    c.add_paragraph("Dòng hai")
    p = tmp_path / "s.docx"
    d.save(p)
    assert docx_lines(p) == ["Dòng một Dòng hai"]
